"""Acrescenta a aba "5. Formas no braco" a planilha da enciclopedia, com as formas de cada acorde em todo o braco.
Uso: python3 tools/enciclopedia/exportar_xlsx.py entrada.xlsx saida.xlsx   (le data/enciclopedia.json)"""
import json, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
ent, sai = sys.argv[1], sys.argv[2]
d = json.load(open("data/enciclopedia.json"))
wb = load_workbook(ent)
if "5. Formas no braço" in wb.sheetnames: del wb["5. Formas no braço"]
ws = wb.create_sheet("5. Formas no braço")
cab = ["Cifra", "Nome do Acorde", "Notas", "Fórmula", "Forma", "Casas (6ª→1ª corda)", "Posição", "Pestana", "Dedos (6ª→1ª corda)"]
ws.append(cab)
for c in d["acordes"]:
    for i, f in enumerate(c["formas"], 1):
        casas = " ".join("x" if x < 0 else str(x) for x in f["v"])
        dedos = " ".join(str(f["dedos"].get(str(6 - k), "-")) if f["v"][k] > 0 else ("0" if f["v"][k] == 0 else "x") for k in range(6))
        pos = "Aberta" if 0 in f["v"] else f"{f['base']}ª casa"
        ws.append([c["cifra"], c["nome"], " · ".join(c["notas"]), " – ".join(c["formula"]), i, casas, pos, "Sim" if f["barre"] else "Não", dedos])
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="14285A"); cell.alignment = Alignment(horizontal="center")
for col, w in zip("ABCDEFGHI", (12, 44, 30, 24, 8, 22, 12, 10, 22)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(sai); print("ok", ws.max_row - 1, "formas")
